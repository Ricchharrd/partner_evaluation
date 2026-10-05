from __future__ import annotations

from io import BytesIO
import re
from typing import Iterable

from .analysis import effective_fact_map, ratio_rows
from .schema import AnalysisProject, STANDARD_ITEMS, facts_to_rows
from .validation import validation_rows
from .periods import period_label


def _literal(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    if isinstance(value, (dict, list)):
        import json
        return json.dumps(value, ensure_ascii=False)
    return value


def _autosize(ws):
    from openpyxl.utils import get_column_letter

    for column in range(1, ws.max_column + 1):
        width = max((len(str(ws.cell(row, column).value or "")) for row in range(1, ws.max_row + 1)), default=10)
        ws.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 12), 45)


def _write_rows(ws, rows: list[dict]):
    from openpyxl.styles import Alignment, Font, PatternFill

    if not rows:
        ws.append(["자료 없음"])
        return
    headers = list(dict.fromkeys(key for row in rows for key in row))
    ws.append(headers)
    for row in rows:
        ws.append([_literal(row.get(header)) for header in headers])
    fill = PatternFill("solid", fgColor="14213D")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autosize(ws)


def build_excel(project: AnalysisProject) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from .storage import public_project_payload

    project = AnalysisProject.from_dict(public_project_payload(project.to_dict()))
    wb = Workbook()
    meta = wb.active
    meta.title = "분석개요"
    meta_rows = [
        ("분석명", project.title), ("프로젝트 ID", project.project_id), ("법인명", project.entity.legal_name),
        ("국가", project.entity.country), ("법인유형", project.entity.entity_type), ("연결/별도", project.entity.reporting_scope),
        ("회계기준", project.entity.accounting_standard), ("상태", project.status), ("최종수정", project.updated_at),
        ("주의", "예비 재무검토 자료이며 정식 신용평가가 아닙니다. 원문과 인적 검토가 필요합니다."),
    ]
    for row in meta_rows:
        meta.append([_literal(value) for value in row])
    meta["A1"].font = Font(bold=True)
    _autosize(meta)

    facts_ws = wb.create_sheet("추출·정규화 데이터")
    _write_rows(facts_ws, facts_to_rows(project.facts))
    ratios_ws = wb.create_sheet("재무비율")
    _write_rows(ratios_ws, ratio_rows(project.ratios))
    validations_ws = wb.create_sheet("검증경고")
    _write_rows(validations_ws, validation_rows(project.validations))
    sources_ws = wb.create_sheet("출처")
    _write_rows(sources_ws, [source.__dict__ for source in project.sources])
    corrections_ws = wb.create_sheet("수정이력")
    _write_rows(
        corrections_ws,
        [
            {
                "fact_id": fact.fact_id, "연도": fact.fiscal_year, "항목": fact.standard_item,
                "원래추출값": fact.normalized_value, "사용자수정값": fact.user_value,
                "수정사유": fact.user_reason, "수정일시": fact.edited_at, "근거": fact.source_locator,
            }
            for fact in project.facts if fact.user_value is not None
        ],
    )
    for name, rows in [
        ("사업역량 근거", project.narrative.get("business_evidence", [])),
        ("기업 동향", project.narrative.get("partner_updates", [])),
        ("작업시간", project.narrative.get("effort_log", [])),
        ("검토 기록", project.narrative.get("audit_log", [])),
        ("환율 근거", project.narrative.get("fx_display", [])),
        ("웹 조사", project.narrative.get("research_briefs", [])),
        ("HITL 검토", project.narrative.get("hitl_review_history", [])),
        ("외부전송 승인", project.narrative.get("hitl_approvals", [])),
    ]:
        _write_rows(wb.create_sheet(name), rows)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _set_cell_text(cell, value):
    cell.text = "-" if value is None else str(value)


def _add_bullets(document, values: Iterable[str]):
    values = list(values or [])
    if not values:
        document.add_paragraph("미확인")
        return
    for value in values:
        document.add_paragraph(str(value), style="List Bullet")


def build_word(project: AnalysisProject) -> bytes:
    from .hitl import current_review
    from .storage import public_project_payload
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    project = AnalysisProject.from_dict(public_project_payload(project.to_dict()))
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    styles = document.styles
    styles["Normal"].font.name = "Malgun Gothic"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    for name, size in (("Title", 15), ("Heading 1", 11), ("Heading 2", 10)):
        style = styles[name]
        style.font.name = "Malgun Gothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)
    title_border = styles["Title"]._element.pPr.find(qn("w:pBdr"))
    if title_border is not None:
        styles["Title"]._element.pPr.remove(title_border)

    title = document.add_paragraph(style="Title")
    title.add_run(f"파트너사 {project.entity.legal_name} 기업 정보")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_line = document.add_paragraph(f"작성 기준  {project.updated_at[:10]}")
    date_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph("공개자료를 바탕으로 회사 기본정보와 재무수치를 정리한 편집용 초안입니다. 확인되지 않은 항목은 원문 대조 후 보완해야 합니다.")

    narrative = project.narrative.get("final") or project.narrative.get("basic") or {}
    document.add_heading("회사 소개", level=1)
    document.add_paragraph(f"회사명  {project.entity.legal_name}")
    document.add_paragraph(f"국가  {project.entity.country or '미확인'}")
    document.add_paragraph(f"업종  {project.entity.industry or '미확인'}")
    if project.entity.official_website:
        document.add_paragraph(f"공식 홈페이지  {project.entity.official_website}")
    overview = narrative.get("company_overview")
    if isinstance(overview, list):
        overview = " ".join(str(value) for value in overview if value)
    if overview and not str(overview).startswith("기업개요는 공식"):
        document.add_paragraph(str(overview))

    document.add_heading("영위 사업", level=1)
    evidence = [row for row in project.narrative.get("business_evidence", []) if row.get("status") == "확인"]
    if evidence:
        for row in evidence[:4]:
            document.add_paragraph(f"{row.get('title', '사업')}  {row.get('description') or row.get('role') or '세부 내용 미확인'}")
    else:
        document.add_paragraph("확인된 사업별 수행 범위 자료가 없습니다. 공식 자료 또는 사내 확인 결과를 기입하십시오.")

    document.add_heading("지분 구조", level=1)
    document.add_paragraph("주요 주주와 지분율은 현재 분석 자료에 포함되지 않았습니다. 공시 또는 공식 연차보고서에서 확인 후 기입하십시오.")

    document.add_heading("재무 현황", level=1)
    selected = effective_fact_map(project.facts)
    years = sorted({year for year, _ in selected})[-3:]
    currencies = {fact.currency for fact in selected.values() if fact.effective_value is not None}
    currency = next(iter(currencies)) if len(currencies) == 1 else ""
    unit = {"EUR": "백만 유로", "USD": "백만 달러", "GBP": "백만 파운드", "KRW": "억 원"}.get(currency)
    divisor = 100_000_000 if currency == "KRW" else 1_000_000
    if not unit:
        unit, divisor = "원문 통화", 1
    scopes = {fact.reporting_scope for fact in selected.values() if fact.effective_value is not None}
    scope = next(iter(scopes)) if len(scopes) == 1 else ("혼합, 확인 필요" if scopes else project.entity.reporting_scope)
    document.add_paragraph(f"단위  {unit}, 범위  {scope}, 회계기준  {project.entity.accounting_standard}")

    def add_financial_table(caption, items):
        document.add_heading(caption, level=2)
        table = document.add_table(rows=1, cols=1 + len(years))
        table.style = "Table Grid"
        _set_cell_text(table.rows[0].cells[0], "항목")
        for index, year in enumerate(years, start=1):
            _set_cell_text(table.rows[0].cells[index], period_label(project.facts, year))
        for item in items:
            cells = table.add_row().cells
            _set_cell_text(cells[0], STANDARD_ITEMS[item])
            for index, year in enumerate(years, start=1):
                fact = selected.get((year, item))
                blocked = any(v.severity == "오류" and v.fiscal_year in (None, year)
                              and (not v.item_keys or item in v.item_keys) for v in project.validations)
                value = fact.effective_value if fact else None
                if blocked:
                    display = "확인 필요"
                elif value is None:
                    display = "미확인"
                elif unit == "원문 통화":
                    display = f"{value:,.0f} {fact.currency}"
                else:
                    display = f"{value / divisor:,.1f}"
                _set_cell_text(cells[index], display)
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.alignment = (WD_ALIGN_PARAGRAPH.LEFT if cell_index == 0
                                           else WD_ALIGN_PARAGRAPH.CENTER)
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        if row_index == 0:
                            run.bold = True

    add_financial_table("손익 및 현금흐름", ["revenue", "operating_income", "net_income", "operating_cash_flow"])
    add_financial_table("자산 부채 및 자본", ["total_assets", "current_assets", "total_liabilities", "current_liabilities", "total_equity"])
    table_note = document.add_paragraph("표의 '미확인'은 0을 뜻하지 않습니다. '확인 필요'는 자동 검증 오류가 있는 값입니다.")
    table_note.paragraph_format.space_before = Pt(6)

    document.add_heading("최근 현안", level=1)
    updates = [row for row in project.narrative.get("partner_updates", []) if row.get("status") == "승인"]
    if updates:
        for row in updates[:4]:
            document.add_paragraph(f"{row.get('published_at', '날짜 미확인')}  {row.get('title', '')}  {row.get('summary', '')}")
    else:
        document.add_paragraph("담당자가 확인한 공개 현안이 이 분석에 포함되지 않았습니다.")

    document.add_heading("자료 범위와 확인사항", level=1)
    document.add_paragraph("이 문서는 공개 재무자료의 편집용 정리본이며 협업 적합성 등급이나 최종 승인이 아닙니다. 재무수치는 보고 법인과 연결·별도 범위에 한정됩니다. 사업부 매출이 제시되어도 사업부의 독립 재무상태나 현금흐름으로 해석하지 않습니다.")
    document.add_paragraph("원문 대조  " + ("담당자 확인 기록 있음" if current_review(project) else "미완료, 원문과 재확인 필요"))
    for issue in project.validations:
        if issue.severity == "오류":
            document.add_paragraph(f"확인 필요  {issue.message}")

    document.add_heading("근거 자료", level=1)
    for source in project.sources:
        document.add_paragraph(f"{source.name}  {source.url or '업로드 문서, 공개 URL 미기록'}")
    document.add_heading("재무수치 원문 위치", level=2)
    for item in ("revenue", "operating_income", "net_income", "operating_cash_flow", "total_assets", "current_assets", "total_liabilities", "current_liabilities", "total_equity"):
        locations = []
        for year in years:
            fact = selected.get((year, item))
            if fact:
                pages = list(dict.fromkeys(re.findall(r"PAGE (\d+)", fact.source_locator)))
                locator = f"PDF {', '.join(pages)}쪽" if pages else (fact.source_locator or "위치 미기록")
                locations.append(f"{year}년 {locator}")
        if locations:
            document.add_paragraph(f"{STANDARD_ITEMS[item]}  {'; '.join(locations)}")
    for row in evidence:
        document.add_paragraph(f"사업 근거  {row.get('title', '')}  {row.get('source_url', '')}  {row.get('locator', '')}")
    for row in updates:
        document.add_paragraph(f"현안 근거  {row.get('title', '')}  {row.get('source_url', '')}")

    output = BytesIO()
    document.save(output)
    return output.getvalue()
