import unittest
from datetime import date
from io import BytesIO
from tempfile import TemporaryDirectory

from openpyxl import load_workbook
from docx import Document

from partner_finance.analysis import calculate_ratios, context_blocks
from partner_finance.intelligence import merge_updates, parse_sec_updates, review_update
from partner_finance.reports import build_excel, build_word
from partner_finance.storage import ProjectStore
from partner_finance.workflow import finalize, is_current, recalculate
from tests.helpers import fact, sample_project


class IntegratedTests(unittest.TestCase):
    def test_context_mixing_and_duplicate_conflicts_block(self):
        for field, value in [("currency", "USD"), ("reporting_scope", "별도"), ("entity_id", "other"), ("period_end", "2024-06-30")]:
            project = sample_project()
            if field == "period_end":
                for f in project.facts:
                    f.period_end = f"{f.fiscal_year}-12-31"
            setattr(project.facts[-1], field, value)
            self.assertIn(2024, context_blocks(project.facts))
            self.assertTrue(all(r.value is None for r in calculate_ratios(project.facts) if r.fiscal_year == 2024))
        project = sample_project()
        project.facts.append(fact(2024, "revenue", 1))
        self.assertIn(2024, context_blocks(project.facts))

    def test_gaps_and_partial_debt(self):
        rows = [fact(2022, "revenue", 100), fact(2024, "revenue", 200), fact(2024, "short_term_debt", 10), fact(2024, "operating_income", 20)]
        ratios = {r.metric_key: r for r in calculate_ratios(rows) if r.fiscal_year == 2024}
        self.assertIsNone(ratios["revenue_growth"].value)
        self.assertIsNone(ratios["debt_to_operating_income"].value)

    def test_corrections_invalidate_legacy_and_narrative(self):
        project = sample_project()
        project.narrative["legacy_company_rating"] = [{"score": 999}]
        recalculate(project)
        self.assertTrue(is_current(project))
        project.narrative["final"] = {"observed_facts": ["stale"]}
        target = next(f for f in project.facts if f.fiscal_year == 2024 and f.standard_item == "interest_expense")
        target.apply_correction(None, "수정 없음")
        target.normalized_value = None
        self.assertFalse(is_current(project))
        recalculate(project)
        self.assertNotIn("final", project.narrative)
        self.assertIsNone(project.narrative["policy_evaluation"][-1]["score"])

    def test_sec_dedup_and_review_does_not_change_scores(self):
        project = sample_project()
        recalculate(project)
        before = project.narrative["policy_evaluation"]
        payload = {"name": "Test", "filings": {"recent": {"filingDate": ["2026-09-01"], "form": ["8-K"], "accessionNumber": ["0000000001-26-000001"], "primaryDocument": ["report.htm"]}}}
        candidates = parse_sec_updates(payload, "1", date(2026, 9, 1), date(2026, 9, 27))
        self.assertEqual(merge_updates(project, candidates), 1)
        identity = project.narrative["partner_updates"][0]["id"]
        review_update(project, identity, True, "Tester", "원문 확인")
        self.assertEqual(merge_updates(project, candidates), 0)
        self.assertEqual(project.narrative["partner_updates"][0]["status"], "승인")
        self.assertEqual(before, project.narrative["policy_evaluation"])

    def test_three_company_roundtrip_and_reports(self):
        with TemporaryDirectory() as root:
            store = ProjectStore(root)
            for i in range(3):
                project = sample_project()
                project.entity.legal_name = f"Synthetic partner {i + 1} (TEST ONLY)"
                recalculate(project)
                finalize(project, "Test reviewer", "합성 입력 통합 테스트. 실기업 평가 아님.")
                store.save(project)
                restored = store.load(project.project_id)
                self.assertTrue(is_current(restored))
                self.assertEqual(len(restored.narrative["assessment_history"]), 1)
                workbook = load_workbook(BytesIO(build_excel(restored)))
                self.assertIn("연도별 평가표", workbook.sheetnames)
                document = Document(BytesIO(build_word(restored)))
                self.assertIn("재무역량 평가표 (잠정)", "\n".join(p.text for p in document.paragraphs))
            self.assertEqual(len(store.list_projects()), 3)

    def test_finalize_requires_current_no_errors_and_owner_isolation(self):
        project = sample_project()
        with self.assertRaises(ValueError):
            finalize(project, "name", "note")
        recalculate(project)
        project.facts[-1].currency = "EUR"
        recalculate(project)
        with self.assertRaises(ValueError):
            finalize(project, "name", "note")
        with TemporaryDirectory() as root:
            store = ProjectStore(root)
            store.save(project, "alice")
            with self.assertRaises(PermissionError):
                store.save(project, "bob")
            with self.assertRaises(KeyError):
                store.load(project.project_id, "bob")

    def test_excel_formula_injection(self):
        project = sample_project()
        project.entity.legal_name = '=HYPERLINK("https://example.com")'
        workbook = load_workbook(BytesIO(build_excel(project)))
        self.assertNotEqual(workbook["분석개요"]["B3"].data_type, "f")


if __name__ == "__main__":
    unittest.main()
