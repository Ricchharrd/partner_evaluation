import unittest
from io import BytesIO

from docx import Document

from partner_finance.reports import build_word
from partner_finance.schema import AnalysisProject, EntityProfile, FinancialFact, SourceDocument


class WordTemplateTests(unittest.TestCase):
    def test_editable_company_brief_keeps_evidence_separate(self):
        entity = EntityProfile("Example Build", country="이탈리아", industry="건설", reporting_scope="연결")
        project = AnalysisProject("Example Build 검토", entity)
        source = SourceDocument("2025 Annual Report", "공개 PDF", url="https://example.com/annual.pdf")
        project.sources.append(source)
        for year, revenue in ((2024, 11_000_000_000), (2025, 12_600_000_000)):
            project.facts.append(FinancialFact(
                entity.entity_id, year, "revenue", "Revenue", revenue / 1000, revenue, "EUR",
                unit_multiplier=1000, period_start=f"{year}-01-01", period_end=f"{year}-12-31",
                source_id=source.source_id, source_locator=f"PAGE {year - 1742}",
            ))
        document = Document(BytesIO(build_word(project)))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_text = [[cell.text for cell in row.cells] for table in document.tables for row in table.rows]
        self.assertIn("파트너사 Example Build 기업 정보", paragraphs)
        self.assertIn("회사 소개", paragraphs)
        self.assertIn("영위 사업", paragraphs)
        self.assertIn("지분 구조", paragraphs)
        self.assertIn("재무 현황", paragraphs)
        self.assertIn("근거 자료", paragraphs)
        self.assertIn(["매출액", "11,000.0", "12,600.0"], table_text)
        self.assertTrue(any("백만 유로" in paragraph for paragraph in paragraphs))
        self.assertTrue(any("지분율은 현재 분석 자료에 포함되지 않았습니다" in paragraph for paragraph in paragraphs))
        self.assertTrue(any("https://example.com/annual.pdf" in paragraph for paragraph in paragraphs))
        self.assertFalse(any("등급 AAA" in paragraph for paragraph in paragraphs))


if __name__ == "__main__":
    unittest.main()
