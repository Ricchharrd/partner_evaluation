from io import BytesIO
from pathlib import Path
import tempfile
import unittest
import zipfile

from openpyxl import load_workbook

from partner_finance.analysis import build_basic_narrative, calculate_ratios
from partner_finance.reports import build_excel, build_word
from partner_finance.storage import ProjectStore
from partner_finance.validation import validate_facts

from tests.helpers import sample_project


class StorageReportTests(unittest.TestCase):
    def test_save_reload_and_source_deduplication(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            store = ProjectStore(tmpdir)
            project = sample_project()
            original_fact_id = project.facts[0].fact_id
            store.save(project, "user-a")
            loaded = store.load(project.project_id, "user-a")
            self.assertEqual(loaded.entity.legal_name, project.entity.legal_name)
            self.assertEqual(loaded.facts[0].fact_id, original_fact_id)
            with self.assertRaises(KeyError):
                store.load(project.project_id, "user-b")
            first, duplicate1 = store.save_source_bytes(project.project_id, "report.csv", b"same")
            second, duplicate2 = store.save_source_bytes(project.project_id, "report.csv", b"same")
            self.assertFalse(duplicate1)
            self.assertTrue(duplicate2)
            self.assertEqual(first.local_path, second.local_path)

    def test_excel_and_word_exports(self):
        project = sample_project()
        project.ratios = calculate_ratios(project.facts)
        project.validations = validate_facts(project.facts)
        project.narrative["final"] = build_basic_narrative(project.facts, project.ratios)
        excel = build_excel(project)
        workbook = load_workbook(BytesIO(excel), read_only=True)
        self.assertIn("추출·정규화 데이터", workbook.sheetnames)
        self.assertIn("검증경고", workbook.sheetnames)
        word = build_word(project)
        with zipfile.ZipFile(BytesIO(word)) as archive:
            xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn("해외 파트너사 기업개요 및 재무평가 보고서", xml)
        self.assertIn("Synthetic Infrastructure Co.", xml)


if __name__ == "__main__":
    unittest.main()
