from __future__ import annotations

import csv
from datetime import date
from io import BytesIO, StringIO
import ipaddress
import mimetypes
from pathlib import Path
import socket
from urllib.parse import urlparse
import urllib.request

from lxml import etree

from .schema import FinancialFact, SourceDocument


MAX_DOWNLOAD_BYTES = 25 * 1024 * 1024
ALLOWED_SCHEMES = {"http", "https"}
ALLOWED_EXTENSIONS = {".pdf", ".xlsx", ".csv", ".xml", ".xhtml", ".html", ".xbrl"}

ITEM_ALIASES = {
    "매출": "revenue", "매출액": "revenue", "revenue": "revenue", "sales": "revenue",
    "영업이익": "operating_income", "operating income": "operating_income", "operating profit": "operating_income",
    "당기순이익": "net_income", "순이익": "net_income", "net income": "net_income", "profit loss": "net_income",
    "자산총계": "total_assets", "총자산": "total_assets", "total assets": "total_assets", "assets": "total_assets",
    "부채총계": "total_liabilities", "총부채": "total_liabilities", "total liabilities": "total_liabilities", "liabilities": "total_liabilities",
    "자본총계": "total_equity", "총자본": "total_equity", "total equity": "total_equity", "equity": "total_equity",
    "유동자산": "current_assets", "current assets": "current_assets",
    "유동부채": "current_liabilities", "current liabilities": "current_liabilities",
    "현금및현금성자산": "cash", "현금": "cash", "cash and cash equivalents": "cash",
    "매출채권": "accounts_receivable", "accounts receivable": "accounts_receivable", "trade receivables": "accounts_receivable",
    "단기차입금": "short_term_debt", "short term debt": "short_term_debt", "short-term debt": "short_term_debt",
    "장기차입금": "long_term_debt", "long term debt": "long_term_debt", "long-term debt": "long_term_debt",
    "금융부채": "financial_debt", "financial debt": "financial_debt", "borrowings": "financial_debt",
    "이자비용": "interest_expense", "interest expense": "interest_expense", "finance costs": "interest_expense",
    "영업활동현금흐름": "operating_cash_flow", "operating cash flow": "operating_cash_flow", "cash flows from operating activities": "operating_cash_flow",
    "투자활동현금흐름": "investing_cash_flow", "investing cash flow": "investing_cash_flow",
    "재무활동현금흐름": "financing_cash_flow", "financing cash flow": "financing_cash_flow",
    "capex": "capex", "capital expenditure": "capex", "자본적지출": "capex",
    "이익잉여금": "retained_earnings", "retained earnings": "retained_earnings",
    "기초현금": "beginning_cash", "beginning cash": "beginning_cash",
    "기말현금": "ending_cash", "ending cash": "ending_cash",
    "환율변동효과": "fx_effect_on_cash", "effect of exchange rates on cash": "fx_effect_on_cash",
}


def normalize_item(value: str) -> str | None:
    cleaned = " ".join(str(value or "").strip().lower().replace("_", " ").split())
    if cleaned in ITEM_ALIASES:
        return ITEM_ALIASES[cleaned]
    collapsed = cleaned.replace(" ", "")
    for alias, item in ITEM_ALIASES.items():
        if alias.replace(" ", "") == collapsed:
            return item
    return None


def parse_number(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace(" ", "")
    if text in {"-", "—", "n/a", "N/A", "미확인"}:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    try:
        number = float(text)
        return -number if negative else number
    except ValueError:
        return None


def _fact_from_row(row: dict, entity_id: str, source: SourceDocument, locator: str) -> FinancialFact | None:
    label = row.get("original_label") or row.get("원문항목") or row.get("item") or row.get("항목") or row.get("표준항목") or ""
    standard_item = row.get("standard_item") or row.get("표준항목코드") or normalize_item(label)
    if not standard_item:
        return None
    year_value = row.get("fiscal_year") or row.get("연도") or row.get("year")
    try:
        fiscal_year = int(float(year_value))
    except (TypeError, ValueError):
        return None
    raw_value = row.get("original_value") if "original_value" in row else row.get("원문값", row.get("value", row.get("값")))
    original_value = parse_number(raw_value)
    multiplier = parse_number(row.get("unit_multiplier") or row.get("단위배수") or 1) or 1.0
    normalized = None if original_value is None else original_value * multiplier
    return FinancialFact(
        entity_id=entity_id,
        fiscal_year=fiscal_year,
        standard_item=str(standard_item),
        original_label=str(label or standard_item),
        original_value=original_value,
        normalized_value=normalized,
        currency=str(row.get("currency") or row.get("통화") or "미확인").upper(),
        unit_multiplier=multiplier,
        period_start=str(row.get("period_start") or row.get("기간시작") or ""),
        period_end=str(row.get("period_end") or row.get("기간종료") or ""),
        reporting_scope=str(row.get("reporting_scope") or row.get("범위") or "연결"),
        accounting_standard=str(row.get("accounting_standard") or row.get("회계기준") or "미확인"),
        source_id=source.source_id,
        source_locator=locator,
        restated=str(row.get("restated") or row.get("재작성") or "").lower() in {"true", "1", "yes", "y", "예"},
        extraction_method="표 구조 파싱",
        validation_status="검토 필요",
    )


def parse_csv(content: bytes, entity_id: str, source: SourceDocument) -> tuple[list[FinancialFact], list[str]]:
    decoded = content.decode("utf-8-sig", errors="replace")
    rows = list(csv.DictReader(StringIO(decoded)))
    facts = []
    for index, row in enumerate(rows, start=2):
        fact = _fact_from_row(row, entity_id, source, f"CSV row {index}")
        if fact:
            facts.append(fact)
    warning = [] if facts else ["표준 long-format 열을 찾지 못했습니다. 샘플 CSV 형식을 사용하거나 수동으로 입력하십시오."]
    return facts, warning


def parse_xlsx(content: bytes, entity_id: str, source: SourceDocument) -> tuple[list[FinancialFact], list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    facts = []
    warnings = []
    for sheet in workbook.worksheets:
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            continue
        headers = [str(value or "").strip() for value in values[0]]
        normalized_headers = {header.lower() for header in headers}
        if {"fiscal_year", "standard_item"}.issubset(normalized_headers) or {"연도", "표준항목코드"}.issubset(set(headers)):
            for row_no, values_row in enumerate(values[1:], start=2):
                row = dict(zip(headers, values_row))
                fact = _fact_from_row(row, entity_id, source, f"{sheet.title}!row {row_no}")
                if fact:
                    facts.append(fact)
            continue

        # Simple wide table: first column is item, remaining headers are fiscal years.
        year_columns = []
        for index, header in enumerate(headers[1:], start=1):
            try:
                year_columns.append((index, int(float(header))))
            except ValueError:
                continue
        if year_columns:
            for row_no, values_row in enumerate(values[1:], start=2):
                label = values_row[0] if values_row else None
                item = normalize_item(label)
                if not item:
                    continue
                for column_index, year in year_columns:
                    value = values_row[column_index] if column_index < len(values_row) else None
                    fact = _fact_from_row(
                        {"fiscal_year": year, "standard_item": item, "original_label": label, "value": value, "currency": "미확인"},
                        entity_id,
                        source,
                        f"{sheet.title}!{sheet.cell(row_no, column_index + 1).coordinate}",
                    )
                    if fact:
                        facts.append(fact)
        else:
            warnings.append(f"시트 '{sheet.title}'에서 지원되는 표 구조를 찾지 못했습니다.")
    return facts, warnings if warnings or facts else ["처리 가능한 재무표를 찾지 못했습니다."]


def extract_pdf_text(content: bytes, max_pages: int = 300) -> tuple[str, list[str]]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content))
    page_count = len(reader.pages)
    pages = [(page.extract_text() or "").strip() for page in reader.pages[:max_pages]]
    text = "\n\n".join(f"[PAGE {index}]\n{page}" for index, page in enumerate(pages, start=1) if page)
    average_chars = sum(len(page) for page in pages) / max(len(pages), 1)
    warnings = []
    if page_count > max_pages:
        warnings.append(f"전체 {page_count}쪽 중 앞 {max_pages}쪽만 텍스트 추출했습니다. 필요한 주석은 별도 파일로 처리하십시오.")
    if average_chars < 80:
        warnings.append("스캔 PDF로 추정됩니다. 이 환경에는 OCR/비전 추출이 연결되어 있지 않아 수동 입력이 필요합니다.")
    else:
        warnings.append("텍스트 PDF를 인식했습니다. 숫자 자동 추출은 Claude API를 명시적으로 실행하거나 표준 CSV/XLSX를 사용하십시오.")
    return text, warnings


XBRL_ITEM_HINTS = {
    "revenue": ["revenue", "salesrevenue", "turnover"],
    "operating_income": ["operatingprofit", "operatingincomeloss"],
    "net_income": ["profitloss", "netincomeloss"],
    "total_assets": ["assets"],
    "total_liabilities": ["liabilities"],
    "total_equity": ["equity", "stockholdersequity"],
    "current_assets": ["currentassets", "assetscurrent"],
    "current_liabilities": ["currentliabilities", "liabilitiescurrent"],
    "cash": ["cashandcashequivalents"],
    "operating_cash_flow": ["cashflowsfromusedinoperatingactivities", "netcashprovidedbyusedinoperatingactivities"],
}


def parse_xbrl(content: bytes, entity_id: str, source: SourceDocument) -> tuple[list[FinancialFact], list[str]]:
    parser = etree.XMLParser(recover=True, huge_tree=False, no_network=True, resolve_entities=False, load_dtd=False)
    root = etree.fromstring(content, parser=parser)
    context_years = {}
    for context in root.xpath("//*[local-name()='context']"):
        context_id = context.get("id")
        end = context.xpath("string(.//*[local-name()='endDate'])") or context.xpath("string(.//*[local-name()='instant'])")
        start = context.xpath("string(.//*[local-name()='startDate'])")
        if context_id and end:
            context_years[context_id] = (int(end[:4]), start, end)

    facts = []
    seen = set()
    for element in root.iter():
        context_ref = element.get("contextRef")
        if not context_ref or context_ref not in context_years or len(element):
            continue
        value = parse_number(element.text)
        if value is None:
            continue
        local = etree.QName(element).localname
        normalized_name = "".join(ch for ch in local.lower() if ch.isalnum())
        item = None
        for key, hints in XBRL_ITEM_HINTS.items():
            if any(normalized_name == hint or normalized_name.endswith(hint) for hint in hints):
                item = key
                break
        if not item:
            continue
        year, start, end = context_years[context_ref]
        dedup_key = (year, item, context_ref)
        if dedup_key in seen:
            continue
        seen.add(dedup_key)
        decimals = element.get("decimals")
        facts.append(
            FinancialFact(
                entity_id=entity_id,
                fiscal_year=year,
                standard_item=item,
                original_label=local,
                original_value=value,
                normalized_value=value,
                currency=(element.get("unitRef") or "미확인").upper(),
                period_start=start,
                period_end=end,
                source_id=source.source_id,
                source_locator=f"XBRL tag {local}, context {context_ref}, decimals {decimals}",
                extraction_method="XBRL/iXBRL 구조 파싱",
                validation_status="검토 필요",
            )
        )
    warning = [] if facts else ["지원되는 XBRL 표준 항목을 찾지 못했습니다. 회사 확장 taxonomy 매핑 또는 수동 검토가 필요합니다."]
    return facts, warning


def parse_uploaded_file(filename: str, content: bytes, entity_id: str, source: SourceDocument):
    if len(content) > MAX_DOWNLOAD_BYTES:
        raise ValueError("파일 크기가 25MB 제한을 초과합니다.")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        facts, warnings = parse_csv(content, entity_id, source)
        return facts, warnings, ""
    if suffix == ".xlsx":
        facts, warnings = parse_xlsx(content, entity_id, source)
        return facts, warnings, ""
    if suffix == ".pdf":
        text, warnings = extract_pdf_text(content)
        return [], warnings, text
    if suffix in {".xml", ".xbrl", ".xhtml", ".html"}:
        facts, warnings = parse_xbrl(content, entity_id, source)
        return facts, warnings, ""
    return [], [f"지원하지 않는 파일 형식입니다: {suffix or '확장자 없음'}"], ""


def validate_public_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme.lower() not in ALLOWED_SCHEMES or not parsed.hostname:
        raise ValueError("http 또는 https 공개 URL만 허용됩니다.")
    if parsed.username or parsed.password:
        raise ValueError("인증정보가 포함된 URL은 허용되지 않습니다.")
    addresses = socket.getaddrinfo(parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80))
    for address in addresses:
        ip = ipaddress.ip_address(address[4][0])
        if not ip.is_global:
            raise ValueError("내부망·로컬·예약 IP 주소에는 접근할 수 없습니다.")


class PublicRedirectHandler(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        validate_public_url(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


def download_public_document(url: str, timeout: int = 20) -> tuple[str, bytes, str]:
    validate_public_url(url)
    request = urllib.request.Request(url, headers={"User-Agent": "Partner Financial Workbench/1.0 contact@example.com"})
    with urllib.request.build_opener(PublicRedirectHandler()).open(request, timeout=timeout) as response:
        final_url = response.geturl()
        validate_public_url(final_url)
        content_type = response.headers.get_content_type()
        content_length = response.headers.get("Content-Length")
        if content_length and int(content_length) > MAX_DOWNLOAD_BYTES:
            raise ValueError("파일 크기가 25MB 제한을 초과합니다.")
        content = response.read(MAX_DOWNLOAD_BYTES + 1)
        if len(content) > MAX_DOWNLOAD_BYTES:
            raise ValueError("파일 크기가 25MB 제한을 초과합니다.")
    name = Path(urlparse(final_url).path).name or "downloaded_document"
    if not Path(name).suffix:
        extension = mimetypes.guess_extension(content_type) or ""
        name += extension
    if Path(name).suffix.lower() not in ALLOWED_EXTENSIONS:
        raise ValueError(f"지원하지 않는 다운로드 형식입니다: {Path(name).suffix or content_type}")
    return name, content, content_type


def sample_csv_bytes() -> bytes:
    rows = [
        ["fiscal_year", "standard_item", "original_label", "value", "currency", "unit_multiplier", "reporting_scope", "period_start", "period_end"],
        [2024, "revenue", "Revenue", 1200, "USD", 1000000, "연결", "2024-01-01", "2024-12-31"],
        [2024, "operating_income", "Operating profit", 90, "USD", 1000000, "연결", "2024-01-01", "2024-12-31"],
        [2024, "total_assets", "Total assets", 2500, "USD", 1000000, "연결", "", "2024-12-31"],
        [2024, "total_liabilities", "Total liabilities", 1500, "USD", 1000000, "연결", "", "2024-12-31"],
        [2024, "total_equity", "Total equity", 1000, "USD", 1000000, "연결", "", "2024-12-31"],
    ]
    buffer = StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")
